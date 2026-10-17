import asyncio
import json
from datetime import datetime
from io import BytesIO
from unittest.mock import AsyncMock
from zipfile import ZIP_DEFLATED, ZipFile

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient

from backend.app.ai import AIError, AIService
from backend.app.config import Settings
from backend.app.demo import DEMO, demo_analysis
from backend.app.documents import extract_document
from backend.app.grounding import assemble_draft, validate_grounding, validate_missing
from backend.app.main import app, ai
from backend.app.schemas import Analysis, DraftSelection, MissingRequest, MissingResult

client = TestClient(app)


def context():
    return {"source": DEMO["source"], "analysis": DEMO["analysis"], "demo": True}


def test_demo_extracts_all_required_scholarship_items_with_exact_evidence():
    response = client.post("/api/analyze", json={"source": DEMO["source"], "demo": True})
    assert response.status_code == 200
    result = response.json()
    assert result["mode"] == "demo"
    analysis = Analysis.model_validate(result["analysis"])
    assert len(analysis.requirements) == 9
    assert {r.id for r in analysis.requirements if r.kind == "document"} == {"certificate", "transcript"}
    assert next(r for r in analysis.requirements if r.kind == "question").character_limit == 500
    assert all(r.evidence in DEMO["source"] for r in analysis.requirements)
    assert validate_grounding(analysis, DEMO["source"]) == analysis


def test_demo_rejects_arbitrary_text_instead_of_reusing_sample_results():
    response = client.post("/api/analyze", json={"source": "Bu metin örnek burs yönergesi değildir. Başka bir ilan.", "demo": True})
    assert response.status_code == 422


def test_fabricated_source_quote_rejected():
    analysis = demo_analysis(DEMO["source"])
    analysis.requirements[0].evidence = "Başvuran kişi en az üç ödül kazanmış olmalıdır."
    with pytest.raises(ValueError, match="birebir"):
        validate_grounding(analysis, DEMO["source"])


def test_fabricated_character_limit_rejected():
    analysis = demo_analysis(DEMO["source"])
    analysis.requirements[4].character_limit = 800
    with pytest.raises(ValueError, match="Karakter"):
        validate_grounding(analysis, DEMO["source"])


def test_invented_deadline_rejected():
    analysis = demo_analysis(DEMO["source"])
    analysis.requirements[6].deadline = datetime.fromisoformat("2026-12-01T23:59:00+03:00")
    with pytest.raises(ValueError, match="tarih"):
        validate_grounding(analysis, DEMO["source"])


def test_model_cannot_assume_timezone():
    analysis = demo_analysis(DEMO["source"])
    analysis.requirements[6].evidence = "Son başvuru tarihi 30 Kasım 2026, saat 23.59'dur."
    with pytest.raises(ValueError, match="Saat dilimi"):
        validate_grounding(analysis, DEMO["source"] + "\n" + analysis.requirements[6].evidence)


def test_duplicate_requirement_ids_and_missing_evidence_rejected():
    data = json.loads(json.dumps(DEMO["analysis"]))
    data["requirements"][1]["id"] = data["requirements"][0]["id"]
    with pytest.raises(ValueError):
        Analysis.model_validate(data)
    data["requirements"][0]["evidence"] = ""
    with pytest.raises(ValueError):
        Analysis.model_validate(data)


def test_demo_missing_starts_with_unconfirmed_eligibility_and_missing_documents():
    response = client.post("/api/missing", json={**context(), "facts": []})
    assert response.status_code == 200
    statuses = {i["requirement_id"]: i["status"] for i in response.json()["items"]}
    assert statuses["student"] == "needs_confirmation"
    assert statuses["certificate"] == "missing"
    assert statuses["transcript"] == "missing"


def test_missing_cannot_claim_unprovided_personal_information():
    req = MissingRequest.model_validate({**context(), "facts": []})
    data = client.post("/api/missing", json={**context(), "facts": []}).json()
    result = MissingResult.model_validate({"items": data["items"]})
    result.items[0].status = "provided"
    result.items[0].user_evidence = "Lisans öğrencisiyim."
    with pytest.raises(ValueError):
        validate_missing(result, req)


def test_missing_must_cover_every_requirement_and_not_invent_na():
    req = MissingRequest.model_validate({**context(), "facts": []})
    data = client.post("/api/missing", json={**context(), "facts": []}).json()
    result = MissingResult.model_validate({"items": data["items"][:-1]})
    with pytest.raises(ValueError):
        validate_missing(result, req)
    result = MissingResult.model_validate({"items": data["items"]})
    result.items[0].status = "not_applicable"
    with pytest.raises(ValueError):
        validate_missing(result, req)


def test_demo_draft_asks_for_personal_information_before_drafting():
    response = client.post("/api/draft", json={**context(), "requirement_id": "motivation", "personal_info": ""})
    assert response.status_code == 200
    assert response.json()["answer"] == ""
    assert len(response.json()["follow_up_questions"]) == 2


def test_demo_draft_only_uses_explicit_personal_information():
    personal = "Yazılım alanında eğitimimi geliştirmek istiyorum.\nBursu eğitim kitapları için kullanacağım."
    response = client.post("/api/draft", json={**context(), "requirement_id": "motivation", "personal_info": personal})
    assert response.status_code == 200
    data = response.json()
    assert data["answer"] == personal
    assert data["mode"] == "demo"
    assert data["character_count"] == len(personal)
    assert data["character_limit"] == 500


def test_draft_rejects_fabricated_achievement_and_over_limit():
    with pytest.raises(ValueError, match="verilmeyen"):
        assemble_draft(DraftSelection(excerpts=["Üç yarışma kazandım."], follow_up_questions=[]), "Öğrenciyim.", 500)
    with pytest.raises(ValueError, match="karakter"):
        assemble_draft(DraftSelection(excerpts=["a" * 501], follow_up_questions=[]), "a" * 501, 500)


def test_followup_blocks_premature_draft():
    result = assemble_draft(DraftSelection(excerpts=["Öğrenciyim."], follow_up_questions=["Hedefiniz nedir?"]), "Öğrenciyim.", 500)
    assert result["answer"] == ""


def test_draft_rejects_unknown_question_and_tampered_analysis():
    data = {**context(), "requirement_id": "unknown", "personal_info": ""}
    assert client.post("/api/draft", json=data).status_code == 422
    data["analysis"] = json.loads(json.dumps(DEMO["analysis"]))
    data["analysis"]["name"] = "Değiştirilmiş demo"
    assert client.post("/api/draft", json=data).status_code == 422


@pytest.mark.parametrize("filename,data", [
    ("bad.exe", b"anything"), ("bad.pdf", b"not pdf"), ("bad.docx", b"not a zip"),
    ("empty.txt", b""), ("large.txt", b"a" * (5 * 1024 * 1024 + 1)), ("binary.txt", b"abc\x00" * 20),
], ids=["unsupported", "fake-pdf", "fake-docx", "empty", "oversized", "binary"])
def test_document_validation(filename, data):
    with pytest.raises(ValueError):
        extract_document(filename, data)


def test_txt_utf8_and_utf16_upload():
    for encoding in ["utf-8-sig", "utf-16"]:
        assert extract_document("sample.txt", DEMO["source"].encode(encoding)) == DEMO["source"]
    response = client.post("/api/documents", files={"file": ("sample.txt", DEMO["source"].encode(), "text/plain")})
    assert response.status_code == 200
    assert response.json()["text"] == DEMO["source"]


def test_docx_paragraphs_and_tables():
    doc = Document()
    doc.add_paragraph("Başvuru için güncel öğrenci belgesi gereklidir.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Transkript belgesi de gereklidir."
    stream = BytesIO()
    doc.save(stream)
    text = extract_document("sample.docx", stream.getvalue())
    assert "öğrenci belgesi" in text and "Transkript" in text


def test_pdf_text_and_scanned_pdf_failure():
    with pymupdf.open() as doc:
        page = doc.new_page()
        page.insert_text((70, 70), "Student certificate and transcript are required for this application.")
        data = doc.tobytes()
    assert "transcript" in extract_document("sample.pdf", data)
    with pymupdf.open() as doc:
        doc.new_page()
        blank = doc.tobytes()
    with pytest.raises(ValueError, match="OCR"):
        extract_document("scan.pdf", blank)


def test_docx_decompression_bomb_limit():
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "a" * (21 * 1024 * 1024))
    with pytest.raises(ValueError, match="sınır"):
        extract_document("bomb.docx", stream.getvalue())


@pytest.mark.parametrize("settings", [Settings(base_url="https://external.example"), Settings(base_url="http://127.0.0.1.evil.test"), Settings(model="llama3"), Settings(model="gemma4:31b-cloud")])
def test_external_endpoints_and_other_models_blocked(settings):
    with pytest.raises(ValueError):
        settings.validate()


def test_status_requires_successful_model_response(monkeypatch):
    service = AIService(Settings())
    service._request = AsyncMock(side_effect=AIError("offline"))
    result = asyncio.run(service.status())
    assert result["state"] == "unavailable"
    service._request = AsyncMock(side_effect=[{"model_info": {"general.architecture": "gemma4"}}, {"message": {"content": '{"ok":true}'}}])
    result = asyncio.run(service.status())
    assert result["state"] == "connected"


def test_status_never_connected_before_first_response_even_at_low_uptime(monkeypatch):
    monkeypatch.setattr("backend.app.ai.time.monotonic", lambda: 10.0)
    service = AIService(Settings())
    service._request = AsyncMock(side_effect=AIError("offline"))
    assert asyncio.run(service.status())["state"] == "unavailable"
    assert service._request.call_count == 1


@pytest.mark.parametrize("show", [[], {"model_info": [] , "details": "wrong"}, {"details": {"families": 3}}])
def test_malformed_model_metadata_is_unavailable_instead_of_crashing(show):
    service = AIService(Settings())
    service._request = AsyncMock(return_value=show)
    assert asyncio.run(service.status())["state"] == "unavailable"


def test_remote_model_on_local_ollama_is_rejected():
    service = AIService(Settings())
    service._request = AsyncMock(return_value={"remote_model": "gemma4", "model_info": {"general.architecture": "gemma4"}})
    with pytest.raises(AIError, match="Bulut"):
        asyncio.run(service.verify_local_model())


def test_live_analysis_uses_configured_gemma4_and_recovers_one_malformed_response(monkeypatch):
    request = AsyncMock(side_effect=[{"model_info": {"general.architecture": "gemma4"}}, {"message": {"content": "malformed"}}, {"message": {"content": json.dumps(DEMO["analysis"])}}])
    monkeypatch.setattr(ai, "_request", request)
    response = client.post("/api/analyze", json={"source": DEMO["source"], "demo": False})
    assert response.status_code == 200
    assert response.json()["mode"] == "live"
    assert request.call_count == 3
    for call in request.call_args_list[1:]:
        assert call.args[2]["model"] == "gemma4:e4b"
        assert call.args[2]["format"]["type"] == "object"


def test_live_failure_does_not_silently_switch_to_demo(monkeypatch):
    monkeypatch.setattr(ai, "_request", AsyncMock(side_effect=AIError("offline")))
    response = client.post("/api/analyze", json={"source": DEMO["source"], "demo": False})
    assert response.status_code == 503
    assert "analysis" not in response.json()


def test_malformed_model_output_not_displayed_after_recovery_fails(monkeypatch):
    monkeypatch.setattr(ai, "_request", AsyncMock(side_effect=[{"model_info": {"general.architecture": "gemma4"}}, {"message": {"content": "wrong"}}, {"message": {"content": "still wrong"}}]))
    response = client.post("/api/analyze", json={"source": DEMO["source"], "demo": False})
    assert response.status_code == 502


def test_input_length_and_external_browser_origin():
    assert client.post("/api/analyze", json={"source": "a" * 40001}).status_code == 422
    assert client.post("/api/analyze", json={"source": DEMO["source"], "demo": True}, headers={"Origin": "https://evil.example"}).status_code == 403
    assert client.post("/api/documents", files={"file": ("large.txt", b"a" * (7 * 1024 * 1024))}).status_code == 413
