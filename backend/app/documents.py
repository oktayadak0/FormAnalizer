from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import pymupdf
from docx import Document

MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_TEXT_LENGTH = 40000


def extract_document(filename: str, data: bytes) -> str:
    if not data or len(data) > MAX_FILE_SIZE:
        raise ValueError("Dosya boş veya 5 MB sınırını aşıyor.")
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".txt":
            if data.startswith((b"\xff\xfe", b"\xfe\xff")):
                text = data.decode("utf-16")
            else:
                text = data.decode("utf-8-sig")
            if "\x00" in text:
                raise ValueError("Geçerli bir metin dosyası seçin.")
        elif suffix == ".pdf":
            if not data.startswith(b"%PDF-"):
                raise ValueError("PDF dosya içeriği geçersiz.")
            with pymupdf.open(stream=data, filetype="pdf") as doc:
                if doc.is_encrypted or doc.page_count > 80:
                    raise ValueError("Şifreli veya 80 sayfadan uzun PDF desteklenmiyor.")
                parts = []
                for page in doc:
                    parts.append(page.get_text())
                    if sum(map(len, parts)) > MAX_TEXT_LENGTH:
                        raise ValueError("Belge 40.000 karakter sınırını aşıyor.")
                text = "\n".join(parts)
        elif suffix == ".docx":
            with ZipFile(BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > 300 or sum(e.file_size for e in entries) > 20 * 1024 * 1024:
                    raise ValueError("DOCX açılmış içerik sınırını aşıyor.")
                if "word/document.xml" not in archive.namelist():
                    raise ValueError("Geçerli bir DOCX dosyası seçin.")
            doc = Document(BytesIO(data))
            text = "\n".join([p.text for p in doc.paragraphs] +
                             [" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
        else:
            raise ValueError("Yalnızca TXT, PDF ve DOCX dosyaları desteklenir.")
    except (UnicodeError, BadZipFile):
        raise ValueError("Dosya okunamadı. TXT için UTF-8 veya UTF-16 kullanın.") from None
    except ValueError:
        raise
    except Exception:
        raise ValueError("Belge bozuk veya desteklenmeyen bir yapıda.") from None
    text = text.strip()
    if len(text) < 30:
        raise ValueError("Yeterli metin bulunamadı. Taranmış PDF için önce cihazınızda OCR uygulayın.")
    if len(text) > MAX_TEXT_LENGTH:
        raise ValueError("Belge 40.000 karakter sınırını aşıyor.")
    return text
